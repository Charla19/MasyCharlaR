"""Render the CV text in the original two-page Arial and navy layout."""

from html import escape
from pathlib import Path
from shutil import copyfile

from docx import Document
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph


ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / "cv" / "generated"
PUBLIC = ROOT / "public"
PAGE_WIDTH, PAGE_HEIGHT = 595.32, 841.92
LEFT, RIGHT = 72, 523
WIDTH = RIGHT - LEFT
NAVY = colors.HexColor("#17365d")
BODY = colors.HexColor("#304057")
MUTED = colors.HexColor("#58677c")
BLUE = colors.HexColor("#0068c8")

pdfmetrics.registerFont(TTFont("Arial", "C:/Windows/Fonts/arial.ttf"))
pdfmetrics.registerFont(TTFont("Arial-Bold", "C:/Windows/Fonts/arialbd.ttf"))
pdfmetrics.registerFont(TTFont("Arial-Italic", "C:/Windows/Fonts/ariali.ttf"))
pdfmetrics.registerFontFamily("Arial", normal="Arial", bold="Arial-Bold", italic="Arial-Italic")


def paragraph(c, text, top, *, x=LEFT, width=WIDTH, size=9.5, leading=11.3, color=BODY, after=0, align=TA_LEFT):
    style = ParagraphStyle("cv", fontName="Arial", fontSize=size, leading=leading, textColor=color, alignment=align)
    item = Paragraph(text, style)
    _, height = item.wrap(width, PAGE_HEIGHT)
    item.drawOn(c, x, PAGE_HEIGHT - top - height)
    return top + height + after


def heading(c, text, top):
    c.saveState()
    letters = c.beginText(LEFT, PAGE_HEIGHT - top - 10)
    letters.setFont("Arial-Bold", 10)
    letters.setFillColor(NAVY)
    letters.setCharSpace(1.15)
    letters.textOut(text)
    c.drawText(letters)
    c.setStrokeColor(NAVY)
    c.setLineWidth(0.8)
    c.line(LEFT, PAGE_HEIGHT - top - 16, RIGHT, PAGE_HEIGHT - top - 16)
    c.restoreState()
    return top + 22


def rich(paragraph_obj):
    chunks = []
    for run in paragraph_obj.runs:
        if not run.text:
            continue
        value = escape(run.text)
        chunks.append(f"<b>{value}</b>" if run.bold else value)
    return "".join(chunks)


def bullet(c, text, top):
    c.setFillColor(colors.black)
    c.circle(80, PAGE_HEIGHT - top - 5, 1.5, stroke=0, fill=1)
    return paragraph(c, escape(text), top, x=87, width=RIGHT - 87, after=2)


def page_one(c, doc):
    p = doc.paragraphs
    c.setFillColor(NAVY)
    c.setFont("Arial-Bold", 17)
    c.drawCentredString(PAGE_WIDTH / 2, PAGE_HEIGHT - 87, p[0].text)
    c.setFont("Arial-Bold", 10.5)
    c.drawCentredString(PAGE_WIDTH / 2, PAGE_HEIGHT - 102, p[1].text)
    contact = escape(p[2].text)
    contact = contact.replace("linkedin.com/in/charla-masy-r", '<link href="https://www.linkedin.com/in/charla-masy-r" color="#0068c8"><u>linkedin.com/in/charla-masy-r</u></link>')
    paragraph(c, contact, 110, width=PAGE_WIDTH - 80, x=40, size=8, leading=10, align=TA_CENTER, color=MUTED)

    top = heading(c, p[4].text, 140)
    top = paragraph(c, escape(p[5].text), top, after=21)
    top = heading(c, p[7].text, top)
    for index in (8, 9, 10, 11):
        top = paragraph(c, rich(p[index]), top, after=2)
    top += 16
    top = heading(c, p[13].text, top)

    for job, first_bullet, last_bullet, tech in ((14, 15, 20, 21), (22, 23, 25, 26), (27, 28, 30, 31)):
        title, organization, date = (run.text for run in p[job].runs[:3])
        markup = f"<b><font color='#17365d'>{escape(title)}</font>{escape(organization)}</b><font color='#58677c'><i>{escape(date)}</i></font>"
        top = paragraph(c, markup, top, size=10, leading=12, after=2)
        for index in range(first_bullet, last_bullet + 1):
            top = bullet(c, p[index].text, top)
        top = paragraph(c, rich(p[tech]), top, size=8.5, leading=10, color=MUTED, after=6)
    if top > 800:
        raise ValueError(f"CV page one overflows: {top:.1f} pt")
    c.showPage()


def page_two(c, doc):
    p = doc.paragraphs
    top = heading(c, p[33].text, 73)
    top = paragraph(c, escape(p[34].text), top, after=3)
    top = paragraph(c, escape(p[35].text), top, after=13)
    top = heading(c, p[36].text, top)
    top = paragraph(c, escape(p[37].text), top, after=21)
    top = heading(c, p[39].text, top)
    top = paragraph(c, f"<b>{escape(p[40].text)}</b>", top, after=3)
    top = paragraph(c, escape(p[41].text), top, x=LEFT + 25, width=WIDTH - 25, after=13)
    top = paragraph(c, f"<b>{escape(p[42].text)}</b>", top, after=3)
    paragraph(c, escape(p[43].text), top, x=LEFT + 25, width=WIDTH - 25)
    c.showPage()


def main():
    for language in ("fr", "en"):
        doc = Document(BUILD / f"MASY-Charla-Rosalie-CV-{language}.docx")
        output = PUBLIC / f"MASY-Charla-Rosalie-CV-{language}.pdf"
        c = canvas.Canvas(str(output), pagesize=(PAGE_WIDTH, PAGE_HEIGHT), pageCompression=1)
        c.setTitle(f"MASY Charla Rosalie | CV ({language.upper()})")
        c.setAuthor("MASY Charla Rosalie")
        page_one(c, doc)
        page_two(c, doc)
        c.save()
        if language == "fr":
            copyfile(output, PUBLIC / "MASY-Charla-Rosalie-CV-Fullstack.pdf")
        print(output)


if __name__ == "__main__":
    main()
